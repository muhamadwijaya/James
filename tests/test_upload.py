import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import json
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch
from PySide6.QtWidgets import QApplication,QLabel
from PySide6.QtCore import Qt,QDate,QTimer
from PySide6.QtTest import QTest
from agregasi.store import Store
from agregasi.settings_runtime import SettingsRuntime
from agregasi.upload_page import SendPage

class UploadTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.store=Store(self.root/'store.sqlite3')
        self.runtime=SettingsRuntime(self.store);self.repo=self.runtime.delivery;self.page=None;self.server=None
        cfg=self.runtime.repo.load();cfg.update(auto_upload=False,retry_count=0);self.runtime.repo.save(cfg)
    def tearDown(self):
        if self.page:self.page.close();self.page.deleteLater()
        self.runtime.shutdown();self.runtime.deleteLater();self.app.processEvents()
        if self.server:self.server.shutdown();self.server.server_close()
        self.store.close();self.temp.cleanup()
    def add(self,n=1,stage='BOX'):
        prefix={'BOX':'BOX','CARTON':'CTN','PALLET':'PLT'}[stage]
        return [self.store.scan(stage,f'{prefix}-UPLOAD-{i:04}-{len(self.store.events(limit=10000))}')[0] for i in range(n)]
    def ui(self):
        self.page=SendPage(self.store);self.page.attach_runtime(self.runtime);return self.page
    def serve(self,handler=None):
        self.requests=[];outer=self
        class Handler(BaseHTTPRequestHandler):
            def log_message(self,*args):pass
            def do_POST(self):
                data=json.loads(self.rfile.read(int(self.headers['Content-Length'])));outer.requests.append((data,self.headers['Idempotency-Key']))
                code,body=handler(data,len(outer.requests)) if handler else (200,{'success':True})
                self.send_response(code);self.end_headers();self.wfile.write(json.dumps(body).encode())
        self.server=ThreadingHTTPServer(('127.0.0.1',0),Handler);threading.Thread(target=self.server.serve_forever,daemon=True).start()
        cfg=self.runtime.repo.load();cfg.update(project_url=f'http://127.0.0.1:{self.server.server_port}',server_host='');self.runtime.repo.save(cfg)
    def wait(self,predicate):
        for _ in range(600):
            self.app.processEvents()
            if predicate():return
            QTest.qWait(10)
        self.fail('Network operation timed out')

    def test_empty_state_never_enqueues_reference_records(self):
        self.assertEqual(self.repo.rows(),[]);p=self.ui();self.assertEqual(p.queue.rowCount(),0);self.assertEqual(p.history.rowCount(),0)
    def test_filters_combine_level_batch_date_and_status(self):
        self.add(3);self.store.configure({'batch':'BATCH-OTHER'});ids=self.add(2,'CARTON')
        self.runtime.repo.mark_delivery(ids[:1],'FAILED');p=self.ui();p.level.setCurrentText('CARTON');p.batch.setCurrentText('BATCH-OTHER');p.filter_status.setCurrentText('FAILED')
        self.assertEqual([r['id'] for r in p.rows],ids[:1]);p.date_enabled.setChecked(True);p.date.setDate(QDate.currentDate().addDays(-1));self.assertEqual(p.rows,[])
    def test_selection_across_pages_and_filter_reset(self):
        self.add(26);p=self.ui();self.assertEqual(p.queue.rowCount(),10);p.queue.item(0,0).setCheckState(Qt.CheckState.Checked)
        chosen=p.rows[0]['id'];p.local_action('page_next');self.assertEqual(p.page_index,1);self.assertIn(chosen,p.selected_ids)
        p.local_action('select_all');self.assertEqual(len(p.selected_ids),26);p.local_action('page_next');self.assertEqual(p.queue.rowCount(),6)
        p.filter_status.setCurrentText('FAILED');self.assertEqual(p.page_index,0);self.assertFalse(p.selected_ids)
    def test_csv_excel_pdf_and_backup_are_valid_exports(self):
        self.add(2);rows=self.repo.rows()
        for kind in ('csv','xlsx','pdf','json'):
            path=self.root/('out.'+kind);self.repo.export(path,rows,kind);self.assertGreater(path.stat().st_size,100)
            if kind=='pdf':self.assertTrue(path.read_bytes().startswith(b'%PDF'))
            else:self.assertEqual(len(self.repo.read_import(path)),2)
        from openpyxl import load_workbook
        wb=load_workbook(self.root/'out.xlsx');self.assertEqual(wb.active.max_row,3);self.assertEqual(wb.active.cell(2,3).value,rows[0]['code']);wb.close()
    def test_export_button_uses_selection_and_real_xlsx(self):
        self.add(12);p=self.ui();p.queue.item(0,0).setCheckState(Qt.CheckState.Checked);path=self.root/'selection.xlsx'
        with patch('agregasi.upload_page.FileDialog.getSaveFileName',return_value=(str(path),'')):
            p.buttons['export_excel'].click()
        self.assertEqual(len(self.repo.read_import(path)),1)
    def test_import_is_deduplicated_and_keeps_record_identity(self):
        self.add(2);path=self.root/'backup.json';self.repo.export(path,self.repo.rows(),'json');records=self.repo.read_import(path)
        self.assertEqual(self.repo.import_records(records),(0,2))
        self.store.db.execute("DELETE FROM outbound_identity");self.store.db.execute("DELETE FROM events WHERE source!='reference'");self.store.db.commit()
        self.assertEqual(self.repo.import_records(records),(2,0));self.assertEqual(self.repo.import_records(records),(0,2))
        self.assertEqual({r['record_id'] for r in self.repo.rows()},{r['record_id'] for r in records})
    def test_invalid_import_does_not_mutate_database(self):
        self.add();path=self.root/'invalid.json';self.repo.export(path,self.repo.rows(),'json');data=json.loads(path.read_text());data['records'].append(data['records'][0]|{'stage':'INVALID','record_id':'bad:id'})
        path.write_text(json.dumps(data));before=len(self.store.events())
        with self.assertRaises(ValueError):self.repo.read_import(path)
        self.assertEqual(len(self.store.events()),before)
    def test_spreadsheet_formula_cells_are_escaped(self):
        self.add();rows=self.repo.rows();rows[0]['note']='=1+1';path=self.root/'safe.xlsx';self.repo.export(path,rows,'xlsx')
        from openpyxl import load_workbook
        wb=load_workbook(path);self.assertNotEqual(wb.active.cell(2,11).data_type,'f');wb.close()
        self.assertEqual(self.repo.read_import(path)[0]['note'],'=1+1')
    def test_selection_upload_does_not_send_other_rows(self):
        ids=self.add(4);self.serve();p=self.ui();p.selected_ids={ids[1]};p.local_action('send_now');self.wait(lambda:not self.runtime.uploading)
        self.assertEqual([r['id'] for r in self.requests[0][0]['events']],[ids[1]])
        self.assertEqual(sum(r['delivery_status']=='PENDING' for r in self.repo.rows()),3)
        self.assertEqual(len(self.repo.history()),1);self.assertGreaterEqual(self.repo.history()[0]['duration_ms'],0)
    def test_all_batches_complete_and_success_is_not_resent(self):
        self.add(5);self.serve();cfg=self.runtime.repo.load();cfg['cache_limit']=2;self.runtime.repo.save(cfg)
        self.runtime.upload();self.wait(lambda:not self.runtime.uploading)
        self.assertEqual([len(r[0]['events']) for r in self.requests],[2,2,1]);self.assertTrue(all(r['delivery_status']=='SUCCESS' for r in self.repo.rows()))
        self.runtime.upload();QTest.qWait(30);self.assertEqual(len(self.requests),3)
    def test_failure_retry_logs_each_attempt_with_same_body_and_key(self):
        self.add();self.serve(lambda data,n:(503,{'error':'busy'}) if n==1 else (200,{'success':True}))
        cfg=self.runtime.repo.load();cfg['retry_count']=1;self.runtime.repo.save(cfg);self.runtime.upload();self.wait(lambda:not self.runtime.uploading)
        self.assertEqual(self.requests[0],self.requests[1]);self.assertEqual([r['delivery_status'] for r in self.repo.history()],['SUCCESS','FAILED'])
        self.assertEqual(self.repo.rows()[0]['attempts'],2)
    def test_pause_during_batch_leaves_remaining_pending(self):
        self.add(3);self.serve();cfg=self.runtime.repo.load();cfg['cache_limit']=1;self.runtime.repo.save(cfg)
        self.runtime.upload();self.runtime.set_paused(True);self.wait(lambda:not self.runtime.uploading)
        self.assertEqual(len(self.requests),1);self.assertEqual(sum(r['delivery_status']=='PENDING' for r in self.repo.rows()),2)
        self.runtime.upload();QTest.qWait(30);self.assertEqual(len(self.requests),1)
        self.runtime.set_paused(False);self.runtime.upload();self.wait(lambda:not self.runtime.uploading);self.assertEqual(len(self.requests),3)
    def test_partial_ack_uses_record_ids_and_unknown_http_200_fails(self):
        ids=self.add(2);self.serve(lambda data,n:(200,{'accepted_record_ids':[data['events'][0]['record_id']]}))
        self.runtime.upload();self.wait(lambda:not self.runtime.uploading)
        self.assertEqual({r['delivery_status'] for r in self.repo.rows()},{'SUCCESS','FAILED'})
        self.assertEqual(len(self.repo.history()),2)
    def test_cache_cleanup_preserves_payload_identity_history_and_pending(self):
        self.add(2);self.serve();first=self.repo.rows()[0];self.runtime.upload(event_ids=[first['id']]);self.wait(lambda:not self.runtime.uploading)
        snapshot=self.repo.snapshot(first);self.assertGreater(self.repo.cache_bytes(),0);before=self.repo.rows();history=self.repo.history()
        self.assertEqual(self.repo.clear_cache(),1);self.assertEqual(self.repo.cache_bytes(),0);self.assertEqual(self.repo.rows(),before);self.assertEqual(self.repo.history(),history);self.assertEqual(self.repo.snapshot(first),snapshot)
    def test_payload_snapshot_survives_edit_and_restart(self):
        self.add();self.serve(lambda data,n:(400,{}));self.runtime.upload();self.wait(lambda:not self.runtime.uploading)
        original=self.requests[0][0]['events'][0];identifier=self.repo.rows()[0]['id'];self.store.db.execute('UPDATE events SET note=? WHERE id=?',('changed later',identifier));self.store.db.commit()
        self.assertEqual(self.repo.snapshot(self.repo.rows()[0]),original)
        self.runtime.shutdown();self.runtime.deleteLater();self.app.processEvents();self.runtime=SettingsRuntime(self.store);self.repo=self.runtime.delivery
        self.assertEqual(self.repo.snapshot(self.repo.rows()[0]),original);self.assertEqual(self.runtime.upload_metrics['http_code'],400)
    def test_auto_sync_setting_is_shared_with_settings(self):
        p=self.ui();p.buttons['toggle_auto'].click();self.assertTrue(self.runtime.repo.load()['auto_upload'])
        p.buttons['toggle_pause'].click();self.assertTrue(self.runtime.paused);self.assertEqual(p.buttons['toggle_pause'].findChild(QLabel,'button_title').text(),'RESUME SYNC')
        with patch.object(self.runtime,'upload') as upload:self.runtime.auto_tick();upload.assert_not_called()

if __name__=='__main__':unittest.main()
