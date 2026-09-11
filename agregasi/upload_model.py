"""Durable outbound identities, attempt history and portable queue exports."""
import csv
import hashlib
import html
import json
import os
import tempfile
from datetime import datetime
from pathlib import Path
from uuid import uuid4

EVENT_FIELDS = ('ts', 'stage', 'action', 'code', 'status', 'note', 'batch', 'product', 'operator', 'source')
EXPORT_FIELDS = ('ts','stage','code','batch','quantity','delivery_status','attempts','record_id','action','status','note','product','operator','source','original_id')
HEADERS = ('WAKTU','LEVEL','KODE','BATCH','JUMLAH','STATUS KIRIM','PERCOBAAN','RECORD ID','AKSI','STATUS DATA','CATATAN','PRODUK','OPERATOR','SUMBER','ID ASAL')


def now(): return datetime.now().isoformat(timespec='seconds')
def safe_cell(value):
    value = str(value)
    return "'"+value if value.lstrip().startswith(('=', '+', '-', '@')) else value


class UploadRepository:
    def __init__(self, store, settings):
        self.store, self.settings = store, settings
        store.db.executescript('''
        CREATE TABLE IF NOT EXISTS outbound_identity(
          event_id INTEGER PRIMARY KEY REFERENCES events(id),record_id TEXT UNIQUE NOT NULL,
          original_id INTEGER NOT NULL,quantity INTEGER NOT NULL DEFAULT 1,snapshot TEXT);
        CREATE TABLE IF NOT EXISTS outbound_cache(request_id TEXT PRIMARY KEY,payload TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS delivery_attempts(
          id INTEGER PRIMARY KEY,event_id INTEGER NOT NULL,ts TEXT NOT NULL,
          status TEXT NOT NULL,http_code INTEGER NOT NULL,duration_ms INTEGER NOT NULL,
          detail TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS delivery_attempts_event ON delivery_attempts(event_id,id);
        ''')
        if not store.get('installation_id'):
            with store.db: store.put('installation_id', str(uuid4()))

    def rows(self):
        rows = self.settings.delivery_rows()
        identities = {r['event_id']:dict(r) for r in self.store.db.execute('SELECT * FROM outbound_identity')}
        quantities = {}
        if self.store.db.execute("SELECT 1 FROM sqlite_master WHERE name='aggregation_runs'").fetchone():
            quantities = {r['parent_code']: r['quantity'] for r in self.store.db.execute('SELECT parent_code,quantity FROM aggregation_runs WHERE parent_code IS NOT NULL')}
        for row in rows:
            identity = identities.get(row['id'], {})
            row['record_id'] = identity.get('record_id', self.store.get('installation_id')+':'+str(row['id']))
            row['quantity'] = identity.get('quantity', max(1,quantities.get(row['code'],1)))
            row['original_id'] = identity.get('original_id',row['id'])
        return rows

    @staticmethod
    def filter_rows(rows, level='', batch='', date='', status=''):
        return [r for r in rows if (not level or r['stage']==level) and
                (not batch or r['batch']==batch) and (not date or r['ts'][:10]==date) and
                (not status or r['delivery_status']==status)]

    def validate_rows(self, rows):
        problems=[]; identities=set()
        for row in rows:
            errors=[]
            if row['stage'] not in ('UNIT','BOX','CARTON','PALLET'): errors.append('level tidak valid')
            if not isinstance(row['code'],str) or not row['code'].strip(): errors.append('kode kosong')
            if not row['batch']: errors.append('batch kosong')
            try: datetime.fromisoformat(row['ts'])
            except (ValueError,TypeError): errors.append('waktu tidak valid')
            if type(row['quantity']) is not int or row['quantity']<1: errors.append('jumlah tidak valid')
            if row['record_id'] in identities: errors.append('record ID duplikat')
            identities.add(row['record_id'])
            if errors: problems.append(f"{row['code'] or row['id']}: {', '.join(errors)}")
        return problems

    def snapshot(self, row):
        old=self.store.db.execute('SELECT snapshot FROM outbound_identity WHERE event_id=?',(row['id'],)).fetchone()
        if old and old['snapshot']: return json.loads(old['snapshot'])
        payload={k:row[k] for k in EVENT_FIELDS}
        payload.update(id=row['original_id'],record_id=row['record_id'],quantity=row['quantity'])
        with self.store.db:
            self.store.db.execute('''INSERT INTO outbound_identity VALUES(?,?,?,?,?)
            ON CONFLICT(event_id) DO UPDATE SET snapshot=excluded.snapshot''',
            (row['id'],row['record_id'],row['original_id'],row['quantity'],json.dumps(payload,ensure_ascii=False)))
        return payload

    def record_attempt(self, ids, accepted, code, duration_ms, detail):
        with self.store.db:
            self.store.db.executemany('INSERT INTO delivery_attempts(event_id,ts,status,http_code,duration_ms,detail) VALUES(?,?,?,?,?,?)',
                [(i,now(),'SUCCESS' if i in accepted else 'FAILED',code,duration_ms,detail) for i in ids])

    def history(self):
        rows={r['id']:r for r in self.rows()}
        result=[]
        for attempt in self.store.db.execute('SELECT * FROM delivery_attempts ORDER BY id DESC'):
            row=rows.get(attempt['event_id'])
            if row: result.append(row | {'history_id':attempt['id'],'delivery_status':attempt['status'],
                'sent_at':attempt['ts'],'duration_ms':attempt['duration_ms'],'http_code':attempt['http_code'],'delivery_detail':attempt['detail']})
        # Preserve pre-v3.6 ledger history without inventing duration or HTTP status.
        seen={r['id'] for r in result}
        for old in self.store.db.execute("SELECT * FROM delivery_ledger WHERE status IN ('SUCCESS','FAILED') ORDER BY updated_at DESC"):
            if old['event_id'] in rows and old['event_id'] not in seen:
                result.append(rows[old['event_id']] | {'history_id':0,'sent_at':old['updated_at'],'duration_ms':None,'http_code':None})
        return sorted(result,key=lambda r:(r['sent_at'],r['history_id']),reverse=True)

    def cache_payload(self,payload):
        with self.store.db:
            self.store.db.execute('INSERT OR REPLACE INTO outbound_cache VALUES(?,?)',
                (payload['request_id'],json.dumps(payload,ensure_ascii=False)))

    def cache_bytes(self):
        return self.store.db.execute('SELECT COALESCE(SUM(length(CAST(payload AS BLOB))),0) FROM outbound_cache').fetchone()[0]

    def clear_cache(self):
        # Payload cache is reproducible; immutable source snapshots and identities remain.
        with self.store.db:
            count=self.store.db.execute('DELETE FROM outbound_cache').rowcount
            self.store.audit('CLEAR CACHE UPLOAD', f'{count} salinan payload; antrean dan riwayat dipertahankan')
        return count

    def export(self, path, rows, kind):
        path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
        fd,tmp=tempfile.mkstemp(prefix='agregasi-',suffix=path.suffix,dir=path.parent);os.close(fd)
        try:
            if kind=='json':
                payload={'application':'AGREGASI-OUTBOX','format_version':1,'created_at':now(),
                         'records':[self.snapshot(r) for r in rows]}
                Path(tmp).write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
            elif kind=='csv':
                with open(tmp,'w',newline='',encoding='utf-8-sig') as f:
                    writer=csv.writer(f);writer.writerow(HEADERS)
                    writer.writerows([[safe_cell(r.get(k,'')) for k in EXPORT_FIELDS] for r in rows])
            elif kind=='xlsx':
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill
                wb=Workbook();ws=wb.active;ws.title='Antrean Upload';ws.append(HEADERS)
                for row in rows:ws.append([row.get(k,0) if k in ('quantity','attempts','original_id') else safe_cell(row.get(k,'')) for k in EXPORT_FIELDS])
                for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='073653')
                from openpyxl.utils import get_column_letter
                for i,w in enumerate((23,12,28,24,12,18,14,52,20,16,42,24,20,18,12),1):ws.column_dimensions[get_column_letter(i)].width=w
                ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;wb.save(tmp);wb.close()
            elif kind=='pdf':
                from PySide6.QtGui import QTextDocument, QPageLayout, QPageSize
                from PySide6.QtPrintSupport import QPrinter
                from PySide6.QtCore import QMarginsF
                printer=QPrinter(QPrinter.PrinterMode.HighResolution)
                printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat);printer.setOutputFileName(tmp)
                printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4));printer.setPageOrientation(QPageLayout.Orientation.Landscape)
                printer.setPageMargins(QMarginsF(12,12,12,12),QPageLayout.Unit.Millimeter)
                body='<h1>AGREGASI — Antrean Upload</h1><p>'+html.escape(now())+f' • {len(rows)} data</p>'
                body+='<table border="1" cellspacing="0" cellpadding="5" width="100%"><thead><tr>'+''.join('<th>'+x+'</th>' for x in HEADERS[:7])+'</tr></thead>'
                body+=''.join('<tr>'+''.join('<td>'+html.escape(str(r.get(k,'')))+'</td>' for k in EXPORT_FIELDS[:7])+'</tr>' for r in rows)+'</table>'
                doc=QTextDocument();doc.setHtml(body);doc.print_(printer)
            else:raise ValueError('Format ekspor tidak didukung.')
            os.replace(tmp,path)
        finally:
            if Path(tmp).exists():Path(tmp).unlink()
        with self.store.db:self.store.audit('EXPORT UPLOAD',f'{kind.upper()}: {len(rows)} data')
        return path

    def read_import(self,path):
        path=Path(path)
        if path.stat().st_size>20*1024*1024:raise ValueError('Berkas impor maksimal 20 MB.')
        if path.suffix.lower()=='.json':
            data=json.loads(path.read_text(encoding='utf-8-sig'))
            if not isinstance(data,dict) or data.get('application')!='AGREGASI-OUTBOX' or data.get('format_version')!=1:
                raise ValueError('Gunakan backup antrean AGREGASI-OUTBOX versi 1.')
            records=data.get('records')
        else:
            if path.suffix.lower()=='.xlsx':
                from zipfile import ZipFile
                with ZipFile(path) as z:
                    if sum(x.file_size for x in z.infolist())>60*1024*1024:raise ValueError('Isi Excel terlalu besar.')
                from openpyxl import load_workbook
                wb=load_workbook(path,read_only=True,data_only=False)
                try:table=list(wb.active.values)
                finally:wb.close()
            elif path.suffix.lower()=='.csv':
                with path.open(encoding='utf-8-sig',newline='') as f:table=list(csv.reader(f))
            else:raise ValueError('Pilih JSON backup, CSV, atau XLSX hasil ekspor Upload.')
            if not table or tuple(table[0])!=HEADERS:raise ValueError('Kolom berkas tidak sesuai ekspor Upload.')
            records=[]
            for values in table[1:]:
                if not any(values):continue
                if len(values)!=len(EXPORT_FIELDS):raise ValueError('Jumlah kolom tidak valid.')
                r=dict(zip(EXPORT_FIELDS,('' if v is None else str(v) for v in values)))
                for key,value in r.items():
                    if value.startswith("'") and value[1:].lstrip().startswith(('=','+','-','@')):r[key]=value[1:]
                try:r['quantity']=int(r['quantity'])
                except ValueError:raise ValueError('Jumlah harus berupa bilangan bulat.')
                # Preserve the original source ID encoded by this application's exports.
                try:r['id']=int(r['original_id'])
                except ValueError:raise ValueError('ID asal harus berupa bilangan bulat.')
                records.append(r)
        if not isinstance(records,list) or len(records)>50000:raise ValueError('Backup harus memuat maksimal 50.000 record.')
        checked=[];seen=set()
        for i,r in enumerate(records,1):
            if not isinstance(r,dict):raise ValueError(f'Record {i} tidak valid.')
            for key in EVENT_FIELDS:
                if key not in r or not isinstance(r[key],str) or len(r[key])>4000:raise ValueError(f'Record {i}: {key} tidak valid.')
            rid=r.get('record_id')
            if not isinstance(rid,str) or not rid.strip() or len(rid)>200 or rid in seen:raise ValueError(f'Record {i}: ID kosong/duplikat.')
            if type(r.get('id')) is not int or r['id']<0:raise ValueError(f'Record {i}: ID asal tidak valid.')
            candidate=r|{'delivery_status':'PENDING'}
            problems=self.validate_rows([candidate])
            if problems:raise ValueError(problems[0])
            seen.add(rid);checked.append({k:r[k] for k in (*EVENT_FIELDS,'id','quantity','record_id')})
        return checked

    def import_records(self, records):
        known={r['record_id'] for r in self.rows()};added=skipped=0
        with self.store.db:
            for record in records:
                if record['record_id'] in known:skipped+=1;continue
                values=[record[k] if k!='source' else 'upload-import' for k in EVENT_FIELDS]
                cur=self.store.db.execute('INSERT INTO events('+','.join(EVENT_FIELDS)+') VALUES('+','.join('?' for _ in EVENT_FIELDS)+')',values)
                self.store.db.execute('INSERT INTO outbound_identity VALUES(?,?,?,?,?)',
                    (cur.lastrowid,record['record_id'],record['id'],record['quantity'],json.dumps(record,ensure_ascii=False)))
                known.add(record['record_id']);added+=1
            self.store.audit('IMPORT ULANG UPLOAD',f'{added} antrean baru; {skipped} ID sudah ada')
        return added,skipped
